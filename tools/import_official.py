"""
Turns the transport office's two spreadsheets into the files the build reads.

    python -I tools/import_official.py "<Transport Route timings.xlsx>" "<student list.xlsx>"

Writes
    data/official/transport.json   the 9 buses exactly as the office gave them:
                                   bus number, driver, stops in order, times.
                                   Committed — it is the source of truth for routes.
    data/private/students.json     every student with the bus and stop their
                                   pickup point was matched to. NOT committed
                                   (gitignored): names and phone numbers of
                                   students do not belong in a public repo.

The office's data is used as given. Where a student's pickup point is spelled
differently from the timing sheet, the reading is recorded next to it; where it
is not on the sheet at all, the student keeps their route and the pickup is left
for the route builder, which looks for it on the road the bus drives.
"""

import datetime
import json
import os
import re
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# Sheet title -> stable id, display name, and the order the office lists them in.
BUSES = [
    ('CHITTOOR', 'chittoor', 'Chittoor', 'KBM Shafi'),
    ('PENUMUR', 'penumur', 'Penumur', 'Yunis Basha'),
    ('CHANDRAGIRI', 'chandragiri', 'Chandragiri', 'A Eswar'),
    ('SRIKALAHASTHI', 'srikalahasti', 'Srikalahasti', 'Manikanta'),
    ('NAGALAPURAM', 'nagalapuram', 'Nagalapuram', 'Yogesh'),
    ('VIJAYAPURAM', 'vijayapuram', 'Vijayapuram', 'Yugandhar'),
    ('TIRUPATHI ROUTE', 'tirupati', 'Tirupati', 'J Shankar'),
    ('SATHRAWADA', 'sathrawada', 'Sathrawada', 'A Antony'),
    ('KVB PURAM', 'kvb-puram', 'KVB Puram', 'V Sudhakar'),
]
SHEET_TO_ID = {b[0]: b[1] for b in BUSES}

# The student list names 8 routes; Nagalapuram is run by two buses.
ROUTE_BUSES = {
    'CHANDRAGIRI ROUTE': ['chandragiri'], 'CHITTOOR ROUTE': ['chittoor'],
    'PENUMUR ROUTE': ['penumur'], 'SATHRAWADA ROUTE': ['sathrawada'],
    'SRIKALAHASTI ROUTE': ['srikalahasti'], 'TIRUPATHI ROUTE': ['tirupati'],
    'VIJAYAPURAM ROUTE': ['vijayapuram'],
    'NAGALAPURAM ROUTE': ['nagalapuram', 'kvb-puram'],
}

# Hand-checked spellings: pickup as written -> stop as on the timing sheet.
ALIAS = {
    'PICHATTUR': 'PICHATUR', 'M R PALLI': 'MR PALLI', 'M R.PALLI': 'MR PALLI',
    'LAKSHMI PURAM': 'LAKSHMIPURAM CIRCLE', 'TIRUCHANUR': 'TIRUCHANOOR',
    'PACHIKAPALLM': 'PACHIKAPALLAM', 'TADUKUPETA': 'THADUKUPETA',
    'NETTERI CROSS': 'NETTIRI CROSS', 'BUGGAAGRAHARAM': 'BUGGA AGRAHARAM',
    'YOGIMALLAVARAM': 'MALLAVARAM', 'GAJULAMANDYAM': 'GAJULAMANDAYAM',
}


def norm(s):
    s = (s or '').upper().replace('.', ' ')
    s = re.sub(r'[^A-Z ]', ' ', s)
    s = re.sub(r'\s+', '', s)
    for a, b in [('TH', 'T'), ('OO', 'U'), ('AA', 'A'), ('EE', 'I'), ('DH', 'D'), ('LL', 'L'),
                 ('TT', 'T'), ('Y', 'I'), ('W', 'V')]:
        s = s.replace(a, b)
    return s


def hhmm(t):
    if isinstance(t, datetime.time):
        return t.strftime('%H:%M')
    if isinstance(t, datetime.datetime):
        return t.strftime('%H:%M')
    m = re.match(r'\s*(\d{1,2})[:.](\d{2})', str(t or ''))
    return f'{int(m[1]):02d}:{m[2]}' if m else None


def read_buses(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {}
    for ws in wb.worksheets:
        key = ws.title.strip()
        if key not in SHEET_TO_ID:
            raise SystemExit(f'unknown sheet "{key}" — add it to BUSES')
        text = ' '.join(str(v) for r in ws.iter_rows(max_row=7, values_only=True) for v in r if v)
        phone = re.findall(r'\b[6-9]\d{9}\b', text)
        reg = re.findall(r'AP\s?\d{2}\s?[A-Z]{1,2}\s?\d{4}', text)
        stops = []
        for r in ws.iter_rows(values_only=True):
            if isinstance(r[0], (int, float)) and len(r) > 1 and r[1]:
                stops.append({'seq': len(stops) + 1, 'sheetNo': int(r[0]),
                              'name': ' '.join(str(r[1]).split()), 'time': hhmm(r[2])})
        out[SHEET_TO_ID[key]] = {'sheet': ws.title.strip(), 'phone': phone[0] if phone else None,
                                 'busNo': reg[0].replace(' ', '') if reg else None, 'stops': stops}
    return out


def read_students(path, buses):
    ws = openpyxl.load_workbook(path, data_only=True).active
    rows = []
    for r in ws.iter_rows(min_row=1, values_only=True):
        if not isinstance(r[0], (int, float)) or not r[1]:
            continue
        rows.append(dict(adm=str(r[1]).strip().upper(), name=' '.join(str(r[2]).split()),
                         cls=' '.join(str(r[3] or '').split()), gender=str(r[5] or '').strip(),
                         mobile=str(r[8] or '').strip(), pickup=' '.join(str(r[9] or '').split()),
                         route=str(r[10] or '').strip().upper()))

    def find(bus, name, loose=False):
        n = norm(ALIAS.get(name.upper(), name))
        for s in buses[bus]['stops']:
            m = norm(s['name'])
            if m == n or (loose and len(n) >= 4 and (m.startswith(n) or n.startswith(m))):
                return s
        return None

    out = []
    for st in rows:
        cands = ROUTE_BUSES.get(st['route'], [])
        exact = [(b, find(b, st['pickup'])) for b in cands]
        exact = [(b, s) for b, s in exact if s]
        loose = [(b, find(b, st['pickup'], True)) for b in cands]
        loose = [(b, s) for b, s in loose if s]
        bus, stop, status, note = None, None, 'MISSING', ''
        if len(exact) == 1:
            bus, stop = exact[0]
            status = 'OK'
            sib = [b for b, _ in loose if b != bus]
            if sib:
                status = 'CHECK'
                note = f'also near a stop on the {sib[0]} bus; placed on {bus}'
            if st['pickup'].upper() in ALIAS:
                note = (note + '; ' if note else '') + f'"{st["pickup"]}" read as "{stop["name"]}"'
        elif not exact and len(loose) == 1:
            bus, stop = loose[0]
            status, note = 'CHECK', f'"{st["pickup"]}" matched to "{stop["name"]}"'
        elif not exact and not loose:
            bus = cands[0] if len(cands) == 1 else None
            note = f'"{st["pickup"]}" is not on the timing sheet'
        else:
            bus, stop = exact[0]
            status, note = 'CHECK', 'matches more than one bus'
        out.append({**st, 'bus': bus, 'candidates': cands, 'stop': stop['name'] if stop else None,
                    'status': status, 'note': note})
    return out


def main():
    timings, students_xlsx = sys.argv[1], sys.argv[2]
    buses = read_buses(timings)
    transport = {
        'year': '2026-27',
        'source': os.path.basename(timings),
        'importedAt': datetime.datetime.now(datetime.timezone.utc).isoformat(timespec='seconds'),
        'buses': [{'id': bid, 'code': str(i + 1), 'name': name, 'sheet': buses[bid]['sheet'],
                   'busNo': buses[bid]['busNo'], 'driver': {'name': drv, 'phone': buses[bid]['phone']},
                   'stops': buses[bid]['stops']}
                  for i, (_, bid, name, drv) in enumerate(BUSES)],
    }
    students = read_students(students_xlsx, buses)

    os.makedirs(os.path.join(ROOT, 'data', 'official'), exist_ok=True)
    os.makedirs(os.path.join(ROOT, 'data', 'private'), exist_ok=True)
    with open(os.path.join(ROOT, 'data', 'official', 'transport.json'), 'w', encoding='utf8') as f:
        json.dump(transport, f, indent=1, ensure_ascii=False)
    with open(os.path.join(ROOT, 'data', 'private', 'students.json'), 'w', encoding='utf8') as f:
        json.dump({'source': os.path.basename(students_xlsx), 'students': students}, f, indent=1, ensure_ascii=False)

    from collections import Counter
    print('buses', len(transport['buses']), 'stops', sum(len(b['stops']) for b in transport['buses']))
    print('students', len(students), dict(Counter(s['status'] for s in students)))
    print('per bus', dict(Counter(s['bus'] for s in students)))


if __name__ == '__main__':
    main()
